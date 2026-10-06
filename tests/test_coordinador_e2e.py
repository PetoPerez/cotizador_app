"""E2E: vendedores cotizan con todas las empresas + rol coordinador de empresa.

Comprueba:
 1. Un vendedor de CLM cotiza con Supliese, y con varias empresas a la vez.
 2. El coordinador no entra a cotizaciones ni clientes (sí al tipo de cambio).
 3. Alta manual del coordinador: queda en su empresa SIN precio y no se cotiza.
 4. Alta de un modelo que ya existe en otra empresa: se liga, no se duplica.
 5. El coordinador edita descripción, pero no precios ni el estado global.
 6. Desactivar desde el coordinador apaga solo su empresa y deja bitácora.
 7. No toca productos que no están en su empresa.
 8. Importación por plantilla sin precios.
 9. Usuarios: un coordinador requiere empresa (y no SDL).

ESCRIBE EN LA BASE DE DATOS: solo corre contra un Postgres local desechable.

    docker run -d --rm --name cotiz_test -e POSTGRES_PASSWORD=test \\
        -e POSTGRES_DB=cotiz -p 55432:5432 postgres:16-alpine
    python tests/test_coordinador_e2e.py
    docker stop cotiz_test
"""
import io
import os
import sys
import uuid

os.environ["DATABASE_URL"] = os.environ.get(
    "TEST_DATABASE_URL", "postgresql://postgres:test@localhost:55432/cotiz")
os.environ.setdefault("SECRET_KEY", "test-secret-key-para-pruebas-locales")
os.environ.setdefault("ALLOWED_ORIGINS", "*")

if not any(h in os.environ["DATABASE_URL"] for h in ("localhost", "127.0.0.1")):
    sys.exit("ABORTADO: este test escribe en la BD y solo corre contra Postgres local. "
             f"DATABASE_URL apunta a: {os.environ['DATABASE_URL'].split('@')[-1]}")

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from fastapi.testclient import TestClient
from app.main import app
from app.database import SessionLocal
from app import models
from app.security import hash_password

client = TestClient(app)
client.__enter__()
TAG = uuid.uuid4().hex[:6]  # modelos/correos únicos: el test se puede repetir en la misma BD
NUM = int(TAG, 16) % 9000 + 500
fallos = []


def check(nombre, cond, detalle=""):
    print(("✅ " if cond else "❌ ") + nombre + ("" if cond else f"  -> {detalle}"))
    if not cond:
        fallos.append(nombre)


def login(email):
    r = client.post("/api/auth/login", json={"email": email, "password": "secret123"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


db = SessionLocal()
emp = {e.codigo: e for e in db.query(models.Empresa).all()}
prod = models.Producto(marca="GIRBAU", equipo="Lavadora", modelo=f"HS-{TAG}")
ajeno = models.Producto(marca="CLM", equipo="Secadora", modelo=f"SC-{TAG}")
db.add_all([prod, ajeno]); db.flush()
db.add_all([
    models.ProductoEmpresa(producto_id=prod.id, empresa_id=emp["clm"].id, precio_lista=100, activo=True),
    models.ProductoEmpresa(producto_id=prod.id, empresa_id=emp["supliese"].id, precio_lista=200, activo=True),
    models.ProductoEmpresa(producto_id=prod.id, empresa_id=emp["supliese_gamesail"].id, precio_lista=300, activo=True),
    models.ProductoEmpresa(producto_id=ajeno.id, empresa_id=emp["clm"].id, precio_lista=50, activo=True),
])
cli = models.Cliente(nombre_razon_social=f"CLIENTE {TAG}")
db.add(cli)


def usuario(rol, codigo, n):
    u = models.Usuario(nombre=f"{rol} {TAG}", email=f"{rol}-{TAG}@test.com",
                       password_hash=hash_password("secret123"), rol=rol,
                       empresa_id=emp[codigo].id, margen_min=-5, margen_max=5, activo=True,
                       numero_corto=NUM + n)  # único: el folio lleva el número de vendedor
    db.add(u)
    return u


usuario("vendedor", "clm", 1)
usuario("coordinador", "supliese_gamesail", 2)
usuario("admin", "clm", 3)
db.commit()
PROD, AJENO, CLI = str(prod.id), str(ajeno.id), str(cli.id)
GS_ID = str(emp["supliese_gamesail"].id)
EMP = {c: str(e.id) for c, e in emp.items()}
db.close()

h_v = login(f"vendedor-{TAG}@test.com")
h_c = login(f"coordinador-{TAG}@test.com")
h_a = login(f"admin-{TAG}@test.com")


def cotizar(h, empresas, producto=PROD):
    return client.post("/api/cotizaciones/", headers=h, json={
        "cliente_id": CLI, "empresas": empresas, "moneda": "USD", "tipo_cambio": 18,
        "items": [{"producto_id": producto, "cantidad": 1, "porcentaje_ajuste": 0}]})


# 1) Vendedor de CLM cotiza con otras empresas
r = cotizar(h_v, ["supliese"])
check("Vendedor CLM cotiza con Supliese (precio 200)",
      r.status_code == 200 and r.json()[0]["items"][0]["precio_lista"] == 200.0, r.text[:200])
r = cotizar(h_v, ["clm", "supliese_gamesail"])
check("Vendedor cotiza con varias empresas a la vez -> 2 cotizaciones",
      r.status_code == 200 and len(r.json()) == 2, r.text[:200])

# 2) Coordinador fuera de cotizaciones/clientes
check("Coordinador no lista cotizaciones (403)", client.get("/api/cotizaciones/", headers=h_c).status_code == 403)
check("Coordinador no cotiza (403)", cotizar(h_c, ["supliese_gamesail"]).status_code == 403)
check("Coordinador no ve clientes (403)", client.get("/api/clientes/", headers=h_c).status_code == 403)
check("Coordinador sí ve el tipo de cambio",
      client.get("/api/cotizaciones/tipo-cambio", headers=h_c).status_code in (200, 503))

# 3) Alta manual: sin precio, en su empresa, no cotizable
r = client.post("/api/productos/", headers=h_c, json={
    "marca": "GAMESAIL", "equipo": "Planchadora", "modelo": f"PL-{TAG}", "descripcion": "nueva",
    "empresas": [{"empresa_id": EMP["clm"], "precio_lista": 1, "activo": True}]})  # se ignora
nuevo = r.json() if r.status_code == 200 else {}
pes = nuevo.get("empresas", [])
check("Alta del coordinador: solo en GS y sin precio",
      len(pes) == 1 and pes[0]["empresa_id"] == GS_ID and pes[0]["precio_lista"] is None, r.text[:200])
r = cotizar(h_v, ["supliese_gamesail"], nuevo.get("id"))
check("Producto sin precio no se puede cotizar (400)", r.status_code == 400, r.text[:200])
lista = client.get("/api/productos/?empresa=supliese_gamesail", headers=h_v).json()
check("Producto sin precio no aparece en el catálogo para cotizar",
      all(p["id"] != nuevo.get("id") for p in lista))

# 4) Modelo existente en otra empresa: se liga, no se duplica
r = client.post("/api/productos/", headers=h_c, json={
    "marca": "clm", "equipo": "secadora", "modelo": f" SC-{TAG} "})
check("Alta de modelo existente lo liga a GS (mismo id, sin duplicar)",
      r.status_code == 200 and r.json()["id"] == AJENO
      and any(pe["empresa_id"] == GS_ID and pe["precio_lista"] is None for pe in r.json()["empresas"]),
      r.text[:200])
r = client.post("/api/productos/", headers=h_c, json={"marca": "clm", "equipo": "secadora", "modelo": f"SC-{TAG}"})
check("Repetir el alta -> 400 'ya está en tu empresa'", r.status_code == 400, r.text[:200])

# 5) Edición: descripción sí, precios/estado global no
r = client.put(f"/api/productos/{PROD}", headers=h_c, json={"descripcion": "Nueva descripción"})
check("Coordinador edita descripción", r.status_code == 200 and r.json()["descripcion"] == "Nueva descripción")
r = client.put(f"/api/productos/{PROD}", headers=h_c, json={
    "empresas": [{"empresa_id": GS_ID, "precio_lista": 1, "activo": True}]})
check("Coordinador NO cambia precios (403)", r.status_code == 403, r.text[:200])
r = client.put(f"/api/productos/{PROD}", headers=h_c, json={"activo": False})
check("Coordinador NO cambia el estado global (403)", r.status_code == 403, r.text[:200])

# 6) Desactivar solo en su empresa
r = client.put(f"/api/productos/{PROD}/activo-empresa", headers=h_c, json={"activo": False})
estados = {pe["empresa_id"]: pe["activo"] for pe in r.json().get("empresas", [])} if r.status_code == 200 else {}
check("Desactivar apaga solo GS; CLM y Supliese siguen activos",
      estados.get(GS_ID) is False and sum(estados.values()) == 2 and r.json()["activo"] is True, r.text[:200])
check("Ya no se cotiza con GS", cotizar(h_v, ["supliese_gamesail"]).status_code == 400)
check("Con CLM se sigue cotizando", cotizar(h_v, ["clm"]).status_code == 200)
hist = client.get(f"/api/productos/{PROD}/historial-estado", headers=h_a).json()
check("La bitácora registra la desactivación en GS",
      hist and hist[0]["empresa_id"] == GS_ID and hist[0]["activo_nuevo"] is False
      and hist[0]["referencia"].endswith("— GS"), str(hist)[:200])
client.put(f"/api/productos/{PROD}/activo-empresa", headers=h_c, json={"activo": True})

# 7) Productos fuera de su empresa
nuevo_clm = client.post("/api/productos/", headers=h_a, json={
    "marca": "CLM", "equipo": "Rodillo", "modelo": f"RD-{TAG}",
    "empresas": [{"empresa_id": EMP["clm"], "precio_lista": 10, "activo": True}]}).json()
r = client.put(f"/api/productos/{nuevo_clm['id']}", headers=h_c, json={"descripcion": "x"})
check("Coordinador no edita productos de otra empresa (403)", r.status_code == 403)
lista_c = client.get("/api/productos/", headers=h_c).json()
check("Coordinador solo lista productos de su empresa",
      lista_c and all(any(pe["empresa_id"] == GS_ID for pe in p["empresas"]) for p in lista_c))

# 8) Importación sin precios
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["marca", "equipo", "modelo", "descripcion"])
ws.append(["GAMESAIL", "Dobladora", f"DB-{TAG}", "importada"])
ws.append(["CLM", "Rodillo", f"RD-{TAG}", ""])  # existe en CLM -> se liga a GS
buf = io.BytesIO(); wb.save(buf)
archivo = {"file": ("p.xlsx", buf.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
r = client.post("/api/productos/importar?confirmar=true&marca_general=true", headers=h_c, files=archivo)
res = r.json().get("resumen", {}) if r.status_code == 200 else {}
check("Importación del coordinador: 1 nuevo + 1 ligado, sin errores",
      res.get("nuevos") == 1 and res.get("actualizar") == 1 and res.get("errores") == 0, r.text[:300])
rd = next((p for p in client.get("/api/productos/", headers=h_a).json() if p["id"] == nuevo_clm["id"]), {})
check("El existente conserva su precio de CLM y entra a GS sin precio",
      {pe["empresa_id"]: pe["precio_lista"] for pe in rd.get("empresas", [])}
      == {EMP["clm"]: 10.0, GS_ID: None}, str(rd)[:300])
r = client.get("/api/productos/plantilla-importar", headers=h_c)
hdr = [c.value for c in next(openpyxl.load_workbook(io.BytesIO(r.content)).active.iter_rows(max_row=1))]
check("La plantilla del coordinador no trae columnas de precio",
      hdr == ["marca", "equipo", "modelo", "descripcion"], str(hdr))

# 9) Usuarios
base = {"nombre": "C", "password": "secret123", "rol": "coordinador"}
r = client.post("/api/usuarios/", headers=h_a, json={**base, "email": f"c1-{TAG}@test.com"})
check("Coordinador sin empresa -> 400", r.status_code == 400, r.text[:200])
r = client.post("/api/usuarios/", headers=h_a, json={
    **base, "email": f"c2-{TAG}@test.com", "empresa_id": EMP["servicios_lavanderia"]})
check("Coordinador de SDL -> 400", r.status_code == 400, r.text[:200])
r = client.post("/api/usuarios/", headers=h_a, json={**base, "email": f"c3-{TAG}@test.com", "empresa_id": GS_ID})
check("Coordinador con empresa de productos -> 200", r.status_code == 200, r.text[:200])

print()
if fallos:
    sys.exit(f"FALLARON {len(fallos)}: {fallos}")
print("OK - test_coordinador_e2e")
